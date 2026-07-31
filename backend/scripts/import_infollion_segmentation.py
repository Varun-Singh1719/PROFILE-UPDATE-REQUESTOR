"""One-time importer: convert the Infollion Research Excel taxonomy into a
CRM Segmentation tree document.

Excel columns (see /tmp/segments.xlsx or the artifact URL):
    ID | Name | Parent ID | Level
    ------------------------------
    Level values in the sheet: L0 / L1 / L2 / L3
    User's mapping (Jul 31 2026):
        L0 rows  -> Level 2 in the segmentation tree
        L1 rows  -> Level 3
        L2 rows  -> Level 4
        L3 rows  -> Level 5
        (Level 1 root is "Infollion Research", injected here)

Tree shape (what /api/segmentations stores):
    {
      "name": "Infollion Research",
      "ext_id": None,
      "children": [
        {"name": "Agriculture", "ext_id": 1, "children": [...]},
        ...
      ]
    }

Behaviour:
    * Overwrites any existing segmentation named "Infollion Research"
      (case-insensitive match).
    * Idempotent: safe to re-run.
    * Actor stamps use the first Super Admin found, or a synthetic
      "importer" stub if none exists yet.
"""
from __future__ import annotations

import asyncio
import os
import sys
import uuid
from pathlib import Path
from typing import Optional

import openpyxl

# Make backend importable when run as `python -m scripts.import_...` or directly.
BACKEND_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACKEND_DIR))

from core import db, now_iso  # noqa: E402


SEGMENT_NAME = "Infollion Research"
DEFAULT_XLSX = "/tmp/segments.xlsx"


def load_rows(xlsx_path: str):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    header = [c.value for c in ws[1]]
    expected = ["ID", "Name", "Parent ID", "Level"]
    if header[:4] != expected:
        raise SystemExit(f"Unexpected header {header!r}, expected {expected!r}")

    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        rows.append({
            "id": int(r[0]),
            "name": (r[1] or "").strip() if r[1] else "",
            "parent_id": int(r[2]) if r[2] is not None else None,
            "level": (r[3] or "").strip() if r[3] else "",
        })
    return rows


def build_tree(rows) -> dict:
    """Return a nested tree rooted at 'Infollion Research'."""
    # Build lookup: id -> node dict
    nodes = {}
    for r in rows:
        nodes[r["id"]] = {
            "name": r["name"],
            "ext_id": r["id"],
            "level": r["level"],          # keep original tag for debugging
            "children": [],
        }

    root_children = []
    for r in rows:
        node = nodes[r["id"]]
        pid = r["parent_id"]
        if pid is None:
            # L0 rows -> children of the injected "Infollion Research" root
            root_children.append(node)
        else:
            parent = nodes.get(pid)
            if parent is None:
                # Orphan (shouldn't happen — we validated upstream). Attach at root.
                root_children.append(node)
            else:
                parent["children"].append(node)

    root = {
        "name": SEGMENT_NAME,
        "ext_id": None,
        "level": "L_ROOT",
        "children": root_children,
    }
    return root


def count_nodes(node: dict) -> int:
    total = 1
    for c in node.get("children", []) or []:
        total += count_nodes(c)
    return total


def depth(node: dict) -> int:
    kids = node.get("children") or []
    if not kids:
        return 1
    return 1 + max(depth(c) for c in kids)


async def pick_actor():
    """Pick the first Super Admin contact for created_by/updated_by stamps."""
    contact = await db["contacts"].find_one(
        {"role": {"$in": ["Super Admin", "SuperAdmin", "super_admin", "Super admin"]}}
    )
    if contact:
        return {
            "id": contact.get("id") or contact.get("_id"),
            "name": contact.get("name"),
            "email": contact.get("email"),
        }
    # Fallback synthetic importer stamp (fresh DB with no users yet)
    return {"id": "system-importer", "name": "System Importer", "email": None}


async def main(xlsx_path: Optional[str] = None):
    xlsx_path = xlsx_path or DEFAULT_XLSX
    if not Path(xlsx_path).exists():
        raise SystemExit(f"Excel file not found at {xlsx_path}")

    print(f"→ Reading {xlsx_path}")
    rows = load_rows(xlsx_path)
    print(f"  loaded {len(rows)} rows")

    by_level = {}
    for r in rows:
        by_level[r["level"]] = by_level.get(r["level"], 0) + 1
    print(f"  by-level counts: {by_level}")

    tree = build_tree(rows)
    print(f"→ Built tree — total nodes: {count_nodes(tree)}, depth: {depth(tree)}")

    actor = await pick_actor()
    print(f"→ Actor for created_by/updated_by: {actor}")

    # Delete any existing "Infollion Research" segmentation(s) (case-insensitive)
    existing = [
        d async for d in db["segmentations"].find(
            {"name": {"$regex": f"^{SEGMENT_NAME}$", "$options": "i"}}
        )
    ]
    if existing:
        ids = [d.get("id") for d in existing]
        print(f"→ Overwriting {len(existing)} existing segmentation(s) with matching name: {ids}")
        await db["segmentations"].delete_many({"id": {"$in": ids}})

    now = now_iso()
    doc = {
        "id": str(uuid.uuid4()),
        "name": SEGMENT_NAME,
        "description": "Imported from Infollion Research taxonomy Excel (L0-L3 → Levels 2-5). Root injected as Level 1.",
        "status": "Active",
        "tree": tree,
        "created_by": actor,
        "created_on": now,
        "updated_by": actor,
        "updated_on": now,
        "source": {
            "kind": "excel_import",
            "file": os.path.basename(xlsx_path),
            "row_count": len(rows),
            "by_level": by_level,
        },
    }
    await db["segmentations"].insert_one(doc)
    print(f"✓ Inserted segmentation {doc['id']} — name={doc['name']!r}")

    # Sanity: re-read
    saved = await db["segmentations"].find_one({"id": doc["id"]})
    saved_tree = saved.get("tree") or {}
    print(f"✓ Verified in DB: total nodes={count_nodes(saved_tree)}, depth={depth(saved_tree)}")

    # Show first few L2 categories
    kids = (saved_tree.get("children") or [])[:5]
    print("  First 5 L2 nodes:")
    for k in kids:
        print(f"    - {k['name']} (ext_id={k['ext_id']}, children={len(k.get('children') or [])})")


if __name__ == "__main__":
    path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    asyncio.run(main(path))
