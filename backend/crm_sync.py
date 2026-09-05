"""CRM sync: external Infollion MySQL (READ-ONLY) → MongoDB.

Two layers:

1. MIRROR — a faithful copy of every MySQL table the DB user has been granted
   SELECT on, stored in Mongo collections prefixed ``mysql_`` (one collection per
   table, ``_id`` = MySQL primary key). Only the granted columns are copied;
   everything else is ignored. Each run takes a full snapshot: rows are
   upserted and rows that disappeared from MySQL are deleted, so the mirror is
   always an exact copy. Derived helper fields (``client_contact_ids``,
   ``client_contact_id``) and indexes make the numbers / filters cheap to
   compute directly in Mongo.

2. NUMBERS — per-client / per-contact metrics computed from the mirror and
   written onto the app's ``clients`` / ``client_contacts`` documents
   (``totals_till_date`` etc.). Definitions (agreed with the business):

   Client
     Contacts  = client contacts CURRENTLY mapped to the client (``client_name``
                 on the Mongo contact; ex-employees live in
                 ``previous_work_experience`` and are NOT counted)
     Projects  = COUNT(DISTINCT projects.id)  WHERE projects.client_id = client
     Serviced  = of those projects, the ones with at least one call having
                 revenue_in_usd > 0
     Calls     = COUNT(DISTINCT calls.id) WHERE calls.fk_project IN (client's projects)
     Revenue   = SUM(calls.revenue_in_usd) over the same calls

   Client Contact — identical, with "the contact's projects" being every
   project whose ``projects.client_contacts`` list contains the contact id.

The app's own record of a MySQL row is kept in ``mysql_ref``:
    mysql_ref = { "mysql_id": <int>, "matched_by": "email"|"name"|"import",
                  "linked_at": iso }

Schedule: twice a day at 15:00 and 23:00 IST (see server.py) + manual buttons.
"""
import io
import uuid
import logging
from collections import defaultdict
from datetime import datetime, date
from decimal import Decimal

from pymongo import ReplaceOne, UpdateOne, ASCENDING

from core import db, now_iso
from mysql_db import mysql_query

logger = logging.getLogger("crm_sync")

CLIENTS = "clients"
CONTACTS = "client_contacts"
SYNC_RUNS = "crm_sync_runs"

CLIENT_TYPES = {
    "Venture Capital/Private Equity",
    "Hedge funds/Public Markets",
    "Research and Consulting",
    "Corporations and Companies",
}

# ----------------------------------------------------------------- mirror spec
# Exactly the columns the MySQL user has SELECT access to (from SHOW GRANTS).
# Tables with full-table SELECT list all their columns explicitly so a new
# (un-granted) column added upstream can never break the sync.
MIRROR_TABLES = {
    "clients": {
        "collection": "mysql_clients",
        "columns": ["id", "name", "type", "contract_valid_till", "created_at", "updated_at",
                    "fk_cem", "client_specific_compliance_requirement", "compliance_start_after",
                    "compliance_end_before", "compliance_email_format", "compliance_description",
                    "meta", "ask_for_govt_employee"],
        "indexes": [("name", ASCENDING), ("type", ASCENDING)],
    },
    "client_contacts": {
        "collection": "mysql_client_contacts",
        "columns": ["id", "salutation", "name", "email", "mobile", "designation", "fkClient",
                    "created_at", "updated_at", "user_id", "is_compliance_officer"],
        "indexes": [("fkClient", ASCENDING), ("email", ASCENDING), ("name", ASCENDING)],
    },
    "client_offices": {
        "collection": "mysql_client_offices",
        "columns": ["id", "name", "entityName", "address", "city", "country", "GSTIN",
                    "client_id", "created_at", "updated_at"],
        "indexes": [("client_id", ASCENDING)],
    },
    "currencies": {
        "collection": "mysql_currencies",
        "columns": ["id", "date", "base_currency", "other_currency", "exchange_rate"],
        "indexes": [("date", ASCENDING)],
    },
    "domains": {
        "collection": "mysql_domains",
        "columns": ["id", "name", "parent_id", "level", "created_at", "updated_at"],
        "indexes": [("parent_id", ASCENDING), ("level", ASCENDING)],
    },
    "projects": {
        "collection": "mysql_projects",
        # column-level grant
        "columns": ["id", "client_id", "client_geography", "l0_domain", "l1_domain", "l2_domain",
                    "l3_domain", "receiving_date", "billing_office", "domain_others",
                    "client_contacts", "type", "category"],
        "indexes": [("client_id", ASCENDING), ("client_contact_ids", ASCENDING),
                    ("receiving_date", ASCENDING), ("type", ASCENDING)],
    },
    "calls": {
        "collection": "mysql_calls",
        # column-level grant
        "columns": ["id", "fk_project", "billing_office_id", "call_start_time", "client_contact",
                    "revenue_in_inr", "revenue_in_usd"],
        "indexes": [("fk_project", ASCENDING), ("client_contact_id", ASCENDING),
                    ("call_start_time", ASCENDING)],
    },
}
MIRROR_BATCH = 2000


# ----------------------------------------------------------------- helpers
def _norm_email(v):
    return (v or "").strip().lower()


def _digits(tok):
    tok = str(tok or "").strip()
    return int(tok) if tok.isdigit() else None


def _id_list(csv):
    """'437,592' → [437, 592] (projects.client_contacts / calls.client_contact)."""
    out = []
    for tok in str(csv or "").split(","):
        n = _digits(tok)
        if n is not None and n not in out:
            out.append(n)
    return out


def _bsonable(v):
    """MySQL → BSON-friendly scalar (Decimal → float, date → datetime)."""
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, date) and not isinstance(v, datetime):
        return datetime(v.year, v.month, v.day)
    return v


async def _alloc_display_ids(counter_key, base_default, n):
    """Reserve `n` sequential display ids from the shared counters collection."""
    if n <= 0:
        return []
    await db["counters"].update_one(
        {"_id": counter_key},
        {"$inc": {"seq": n}, "$setOnInsert": {"seq_base": base_default}},
        upsert=True,
    )
    doc = await db["counters"].find_one({"_id": counter_key})
    base = int((doc or {}).get("seq_base") or base_default)
    seq = int((doc or {}).get("seq") or n)
    start = base + seq - n + 1
    return list(range(start, start + n))


# ----------------------------------------------------------------- mirror
async def mirror_table(table: str) -> dict:
    """Snapshot one MySQL table into its ``mysql_*`` Mongo collection."""
    spec = MIRROR_TABLES[table]
    coll = db[spec["collection"]]
    cols = ", ".join(f"`{c}`" for c in spec["columns"])
    rows = await mysql_query(f"SELECT {cols} FROM `{table}`")
    synced_at = now_iso()

    ops, seen = [], set()
    for r in rows:
        doc = {k: _bsonable(r.get(k)) for k in spec["columns"]}
        doc["_id"] = r["id"]
        doc["_synced_at"] = synced_at
        if table == "projects":
            doc["client_contact_ids"] = _id_list(r.get("client_contacts"))
        elif table == "calls":
            doc["client_contact_id"] = _digits(r.get("client_contact"))
            doc["revenue_in_usd"] = float(r.get("revenue_in_usd") or 0)
            doc["revenue_in_inr"] = float(r.get("revenue_in_inr") or 0)
        seen.add(r["id"])
        ops.append(ReplaceOne({"_id": r["id"]}, doc, upsert=True))

    upserted = 0
    for i in range(0, len(ops), MIRROR_BATCH):
        res = await coll.bulk_write(ops[i:i + MIRROR_BATCH], ordered=False)
        upserted += (res.upserted_count or 0)
    # Mirror deletions: anything not in this snapshot is gone upstream.
    deleted = 0
    if seen:
        res = await coll.delete_many({"_id": {"$nin": list(seen)}})
        deleted = res.deleted_count
    else:
        res = await coll.delete_many({})
        deleted = res.deleted_count
    for field, direction in spec["indexes"]:
        try:
            await coll.create_index([(field, direction)])
        except Exception:  # noqa: BLE001 — index creation is best-effort
            pass
    stats = {"rows": len(rows), "inserted": upserted, "deleted": deleted}
    logger.info(f"mirror {table} → {spec['collection']}: {stats}")
    return stats


async def mirror_all() -> dict:
    out = {}
    for table in MIRROR_TABLES:
        out[table] = await mirror_table(table)
    return out


# ----------------------------------------------------------------- metrics (from the mirror)
async def _load_projects_and_calls():
    projects = [p async for p in db["mysql_projects"].find(
        {}, {"_id": 1, "client_id": 1, "client_contact_ids": 1, "receiving_date": 1})]
    calls = [c async for c in db["mysql_calls"].find(
        {}, {"_id": 1, "fk_project": 1, "call_start_time": 1, "revenue_in_usd": 1})]
    return projects, calls


def _metrics_from(pids: set, calls_by_project: dict, proj_recv: dict, proj_has_rev: set):
    """Numbers for one entity given the set of ITS project ids."""
    call_ids, rev, last_call = set(), 0.0, None
    for pid in pids:
        for c in calls_by_project.get(pid, ()):
            call_ids.add(c["_id"])
            rev += float(c.get("revenue_in_usd") or 0)
            st = c.get("call_start_time")
            if st and (last_call is None or st > last_call):
                last_call = st
    recv = [proj_recv[p] for p in pids if proj_recv.get(p)]
    return {
        "totals_till_date": {
            "projects": len(pids),                      # COUNT(DISTINCT project id)
            "serviced": len(pids & proj_has_rev),       # projects with a paid call
            "calls": len(call_ids),                     # COUNT(DISTINCT call id) via fk_project
            "revenue": int(round(rev)),                 # SUM(revenue_in_usd) of those calls
        },
        "last_project_receiving_date": max(recv) if recv else None,
        "last_call_date": last_call,
    }


async def compute_all_metrics():
    """Per-contact and per-client metrics keyed by MySQL id (from the mirror)."""
    projects, calls = await _load_projects_and_calls()

    proj_recv, client_projects, contact_projects = {}, defaultdict(set), defaultdict(set)
    for p in projects:
        pid = p["_id"]
        proj_recv[pid] = p.get("receiving_date")
        if p.get("client_id") is not None:
            client_projects[p["client_id"]].add(pid)
        for cid in p.get("client_contact_ids") or []:
            contact_projects[cid].add(pid)

    calls_by_project, proj_has_rev = defaultdict(list), set()
    for c in calls:
        calls_by_project[c.get("fk_project")].append(c)
        if float(c.get("revenue_in_usd") or 0) > 0:
            proj_has_rev.add(c.get("fk_project"))

    contact_metrics = {cid: _metrics_from(pids, calls_by_project, proj_recv, proj_has_rev)
                       for cid, pids in contact_projects.items()}
    client_metrics = {}
    for clid, pids in client_projects.items():
        m = _metrics_from(pids, calls_by_project, proj_recv, proj_has_rev)
        m["project_count"] = m["totals_till_date"]["projects"]
        m["serviced_count"] = m["totals_till_date"]["serviced"]
        client_metrics[clid] = m
    return contact_metrics, client_metrics


async def _metrics_for_projects(pids: set):
    if not pids:
        return _metrics_from(set(), {}, {}, set())
    pid_list = list(pids)
    proj_recv = {p["_id"]: p.get("receiving_date") async for p in db["mysql_projects"].find(
        {"_id": {"$in": pid_list}}, {"_id": 1, "receiving_date": 1})}
    calls_by_project, proj_has_rev = defaultdict(list), set()
    async for c in db["mysql_calls"].find({"fk_project": {"$in": pid_list}},
                                          {"_id": 1, "fk_project": 1, "call_start_time": 1, "revenue_in_usd": 1}):
        calls_by_project[c["fk_project"]].append(c)
        if float(c.get("revenue_in_usd") or 0) > 0:
            proj_has_rev.add(c["fk_project"])
    return _metrics_from(pids, calls_by_project, proj_recv, proj_has_rev)


async def contact_metrics_for(mysql_id):
    """Numbers for ONE contact (per-contact Sync button) — from the mirror."""
    cid = int(mysql_id)
    pids = {p["_id"] async for p in db["mysql_projects"].find({"client_contact_ids": cid}, {"_id": 1})}
    return await _metrics_for_projects(pids)


async def client_metrics_for(mysql_id):
    """Numbers for ONE client (per-client Sync button) — from the mirror."""
    clid = int(mysql_id)
    pids = {p["_id"] async for p in db["mysql_projects"].find({"client_id": clid}, {"_id": 1})}
    m = await _metrics_for_projects(pids)
    m["project_count"] = m["totals_till_date"]["projects"]
    m["serviced_count"] = m["totals_till_date"]["serviced"]
    return m


# ----------------------------------------------------------------- sync: clients
async def sync_clients(client_metrics=None):
    """Write client numbers onto app `clients` (match by mysql_ref, then by
    case-insensitive name; create missing). Bulk — one round-trip per 1k rows."""
    stats = defaultdict(int)
    if client_metrics is None:
        _, client_metrics = await compute_all_metrics()

    rows = [r async for r in db["mysql_clients"].find({}, {"_id": 1, "name": 1, "type": 1, "created_at": 1})]
    by_mid, by_name = {}, {}
    async for c in db[CLIENTS].find({}, {"_id": 1, "name": 1, "mysql_ref.mysql_id": 1}):
        mid = (c.get("mysql_ref") or {}).get("mysql_id")
        if mid is not None:
            by_mid.setdefault(mid, c["_id"])
        by_name.setdefault((c.get("name") or "").strip().lower(), c["_id"])

    ops, to_create, ts = [], [], now_iso()
    for r in rows:
        mid = r["_id"]
        name = (r.get("name") or "").strip()
        ctype = r.get("type") if r.get("type") in CLIENT_TYPES else (r.get("type") or "")
        metrics = client_metrics.get(mid, {})
        totals = metrics.get("totals_till_date", {"projects": 0, "serviced": 0, "calls": 0, "revenue": 0})
        set_fields = {
            "mysql_ref": {"mysql_id": mid, "matched_by": "name", "linked_at": ts},
            "totals_till_date": totals,
            "project_count": metrics.get("project_count", 0),
            "serviced_count": metrics.get("serviced_count", 0),
            "last_project_receiving_date": metrics.get("last_project_receiving_date"),
            "last_call_date": metrics.get("last_call_date"),
            "updated_on": ts,
        }
        target = by_mid.get(mid) or by_name.get(name.lower())
        if target is not None:
            ops.append(UpdateOne({"_id": target}, {"$set": set_fields}))
            stats["matched_linked"] += 1
        else:
            to_create.append((r, name, ctype, set_fields))

    for i in range(0, len(ops), 1000):
        await db[CLIENTS].bulk_write(ops[i:i + 1000], ordered=False)

    if to_create:
        ids = await _alloc_display_ids("client_seq", 1000, len(to_create))
        docs = []
        for (r, name, ctype, set_fields), disp in zip(to_create, ids):
            created = r.get("created_at")
            docs.append({
                "id": str(uuid.uuid4()),
                "display_id": disp,
                "name": name,
                "type": ctype,
                "key_account_manager_ids": [],
                "client_contact_count": 0,
                "created_by": {"name": "MySQL Sync"},
                "created_on": created.isoformat() if isinstance(created, datetime) else (created or ts),
                "updated_by": {"name": "MySQL Sync"},
                **set_fields,
            })
        if docs:
            await db[CLIENTS].insert_many(docs)
            stats["created"] += len(docs)
    logger.info(f"sync_clients: {dict(stats)}")
    return dict(stats)


# ----------------------------------------------------------------- sync: contacts
async def sync_contacts(contact_metrics=None):
    """Write contact numbers onto app `client_contacts` (match by mysql_ref,
    then by case-insensitive email; create missing). Bulk."""
    stats = defaultdict(int)
    if contact_metrics is None:
        contact_metrics, _ = await compute_all_metrics()

    client_name_by_id = {c["_id"]: (c.get("name") or "").strip()
                         async for c in db["mysql_clients"].find({}, {"_id": 1, "name": 1})}

    rows = [r async for r in db["mysql_client_contacts"].find(
        {}, {"_id": 1, "salutation": 1, "name": 1, "email": 1, "mobile": 1, "designation": 1, "fkClient": 1})]
    by_mid, by_email = {}, {}
    async for c in db[CONTACTS].find({}, {"_id": 1, "email": 1, "mysql_ref.mysql_id": 1}):
        mid = (c.get("mysql_ref") or {}).get("mysql_id")
        if mid is not None:
            by_mid.setdefault(mid, c["_id"])
        em = _norm_email(c.get("email"))
        if em:
            by_email.setdefault(em, c["_id"])

    empty = {"projects": 0, "serviced": 0, "calls": 0, "revenue": 0}
    ops, to_create, ts = [], [], now_iso()
    for r in rows:
        mid = r["_id"]
        email = _norm_email(r.get("email"))
        metrics = contact_metrics.get(mid, {})
        set_fields = {
            "mysql_ref": {"mysql_id": mid, "matched_by": "email", "linked_at": ts},
            "totals_till_date": metrics.get("totals_till_date", empty),
            "last_project_receiving_date": metrics.get("last_project_receiving_date"),
            "last_call_date": metrics.get("last_call_date"),
            "updated_on": ts,
        }
        target = by_mid.get(mid) or (by_email.get(email) if email else None)
        if target is not None:
            ops.append(UpdateOne({"_id": target}, {"$set": set_fields}))
            stats["matched_linked"] += 1
        else:
            to_create.append((r, email, set_fields))

    for i in range(0, len(ops), 1000):
        await db[CONTACTS].bulk_write(ops[i:i + 1000], ordered=False)

    if to_create:
        ids = await _alloc_display_ids("client_contact_seq", 1041, len(to_create))
        docs = []
        for (r, email, set_fields), disp in zip(to_create, ids):
            docs.append({
                "id": str(uuid.uuid4()),
                "display_id": disp,
                "name": (r.get("name") or "").strip(),
                "email": (r.get("email") or "").strip(),
                "phone": (r.get("mobile") or "").strip(),
                "phone_isd": "",
                "client_name": client_name_by_id.get(r.get("fkClient"), ""),
                "designation": (r.get("designation") or "").strip(),
                "type": "",
                "base_location": "",
                "city": "",
                "country_id": None,
                "country_name": "",
                "industries": [],
                "previous_work_experience": [],
                "linkedin_url": "",
                "salutation": (r.get("salutation") or "").strip(),
                "status": "Active",
                "created_by": {"name": "MySQL Sync"},
                "created_on": ts,
                "updated_by": {"name": "MySQL Sync"},
                **set_fields,
            })
        if docs:
            await db[CONTACTS].insert_many(docs)
            stats["created"] += len(docs)
    logger.info(f"sync_contacts: {dict(stats)}")
    return dict(stats)


# ----------------------------------------------------------------- contacts per client
async def refresh_client_contact_counts():
    """`client_contact_count` = contacts CURRENTLY mapped to the client
    (Mongo client_contacts grouped by `client_name`; ex-employees are in
    `previous_work_experience` and are not counted). Clients with none → 0."""
    counts = {}
    cursor = db[CONTACTS].aggregate([
        {"$match": {"client_name": {"$nin": [None, ""]}}},
        {"$group": {"_id": "$client_name", "n": {"$sum": 1}}},
    ])
    async for g in cursor:
        counts[g["_id"]] = g["n"]
    await db[CLIENTS].update_many(
        {"name": {"$nin": list(counts.keys())}}, {"$set": {"client_contact_count": 0}}
    )
    ops = [UpdateOne({"name": name}, {"$set": {"client_contact_count": n}}) for name, n in counts.items()]
    if ops:
        await db[CLIENTS].bulk_write(ops, ordered=False)
    logger.info(f"refresh_client_contact_counts: {len(counts)} clients with contacts")
    return len(counts)


# ----------------------------------------------------------------- orchestration
async def run_full_sync(scope="all", trigger="manual"):
    """1) mirror MySQL → Mongo, 2) recompute numbers, 3) write onto app records.
    Every run is logged in `crm_sync_runs` (latest = "last synced")."""
    started = now_iso()
    run = {"id": str(uuid.uuid4()), "scope": scope, "trigger": trigger,
           "started_at": started, "status": "running"}
    await db[SYNC_RUNS].insert_one(dict(run))
    result = {}
    try:
        result["mirror"] = await mirror_all()
        contact_metrics, client_metrics = await compute_all_metrics()
        if scope in ("all", "clients"):
            result["clients"] = await sync_clients(client_metrics)
        if scope in ("all", "contacts", "client-contacts"):
            result["contacts"] = await sync_contacts(contact_metrics)
        result["client_contact_counts"] = await refresh_client_contact_counts()
        result["ok"] = True
    except Exception as e:  # noqa: BLE001
        logger.exception("run_full_sync failed")
        result["ok"] = False
        result["error"] = str(e)
    await db[SYNC_RUNS].update_one(
        {"id": run["id"]},
        {"$set": {"status": "ok" if result.get("ok") else "failed",
                  "finished_at": now_iso(), "result": result}},
    )
    return result


async def last_sync_run():
    doc = await db[SYNC_RUNS].find_one({}, {"_id": 0}, sort=[("started_at", -1)])
    return doc


# ----------------------------------------------------------------- Excel export of the mirror
async def build_mysql_excel() -> bytes:
    """One workbook, one tab per mirrored MySQL table (granted columns only)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="EC9324")
    for table, spec in MIRROR_TABLES.items():
        ws = wb.create_sheet(title=table[:31])
        cols = spec["columns"]
        ws.append(cols)
        for i in range(1, len(cols) + 1):
            c = ws.cell(row=1, column=i)
            c.font, c.fill = head_font, head_fill
        n = 0
        async for r in db[spec["collection"]].find({}, {c: 1 for c in cols}).sort("_id", 1):
            row = []
            for c in cols:
                v = r.get(c)
                if isinstance(v, (dict, list)):
                    v = str(v)
                row.append(v)
            ws.append(row)
            n += 1
        for i, c in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(c) + 4))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        logger.info(f"excel: {table} → {n} rows")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
